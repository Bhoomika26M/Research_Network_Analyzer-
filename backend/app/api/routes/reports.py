"""Reports with filters: JSON by default, XLSX (openpyxl) or PDF (reportlab) export."""

import io
import uuid
from typing import Any, Literal

from fastapi import APIRouter, Depends, Query
from fastapi.encoders import jsonable_encoder
from fastapi.responses import JSONResponse, Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user, get_db
from app.models.enums import PublicationStatus, PublicationType
from app.schemas.report import ReportEnvelope
from app.services import reports as report_service

router = APIRouter(
    prefix="/api/v1/reports",
    tags=["reports"],
    dependencies=[Depends(get_current_user)],
)

Format = Literal["json", "xlsx", "pdf"]
Column = tuple[str, str]


def _xlsx_bytes(title: str, columns: list[Column], rows: list[dict[str, Any]]) -> bytes:
    """Render rows as a single-sheet .xlsx workbook (imported lazily)."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "Report"
    sheet.append([header for _, header in columns])
    for row in rows:
        sheet.append([row.get(key, "") for key, _ in columns])
    for index, (_, header) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 4)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(title: str, columns: list[Column], rows: list[dict[str, Any]]) -> bytes:
    """Render rows as a simple PDF table (reportlab imported lazily)."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements: list[Any] = [Paragraph(title.title() + " report", styles["Title"]), Spacer(1, 12)]
    data = [[header for _, header in columns]]
    for row in rows:
        data.append([str(row.get(key, "")) for key, _ in columns])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), "#1f2937"),
                ("TEXTCOLOR", (0, 0), (-1, 0), "#ffffff"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, "#9ca3af"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buffer.getvalue()


def _render(
    report: str,
    columns: list[Column],
    rows: list[dict[str, Any]],
    fmt: Format,
) -> Response:
    if fmt == "json":
        envelope = ReportEnvelope(
            report=report,
            count=len(rows),
            columns=[header for _, header in columns],
            rows=rows,
        )
        return JSONResponse(jsonable_encoder(envelope))
    if fmt == "xlsx":
        body = _xlsx_bytes(report, columns, rows)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        headers = {"Content-Disposition": f'attachment; filename="{report}-report.xlsx"'}
    else:
        body = _pdf_bytes(report, columns, rows)
        media = "application/pdf"
        headers = {"Content-Disposition": f'attachment; filename="{report}-report.pdf"'}
    return Response(content=body, media_type=media, headers=headers)


@router.get("/publications")
async def publications_report(
    db: AsyncSession = Depends(get_db),
    format: Format = "json",
    year_from: int | None = Query(default=None, ge=1, le=9999),
    year_to: int | None = Query(default=None, ge=1, le=9999),
    type: PublicationType | None = None,
    status: PublicationStatus | None = None,
    institution_id: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
) -> Response:
    parsed_institution = uuid.UUID(institution_id) if institution_id else None
    columns, rows = await report_service.publications_report(
        db,
        year_from=year_from,
        year_to=year_to,
        type_=type,
        status=status,
        institution_id=parsed_institution,
        limit=limit,
    )
    return _render("publications", columns, rows, format)


@router.get("/research")
async def research_report(
    db: AsyncSession = Depends(get_db),
    format: Format = "json",
    institution_id: str | None = None,
    min_publications: int = Query(0, ge=0),
    limit: int = Query(500, ge=1, le=5000),
) -> Response:
    parsed_institution = uuid.UUID(institution_id) if institution_id else None
    columns, rows = await report_service.research_report(
        db,
        institution_id=parsed_institution,
        min_publications=min_publications,
        limit=limit,
    )
    return _render("research", columns, rows, format)


@router.get("/collaboration")
async def collaboration_report(
    db: AsyncSession = Depends(get_db),
    format: Format = "json",
    status: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
) -> Response:
    columns, rows = await report_service.collaboration_report(db, status=status, limit=limit)
    return _render("collaboration", columns, rows, format)


@router.get("/institutions")
async def institutions_report(
    db: AsyncSession = Depends(get_db),
    format: Format = "json",
    country: str | None = None,
    limit: int = Query(500, ge=1, le=5000),
) -> Response:
    columns, rows = await report_service.institutions_report(db, country=country, limit=limit)
    return _render("institutions", columns, rows, format)
